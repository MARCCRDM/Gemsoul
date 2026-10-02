from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path
out=Path('docs/boss-design'); d=Document(); s=d.sections[0]
s.top_margin=s.bottom_margin=Inches(.7); s.left_margin=s.right_margin=Inches(.75)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
 st=d.styles[name]; st.font.name='Calibri'; st.font.color.rgb=RGBColor(0,0,0)
d.styles['Normal'].font.size=Pt(11); d.styles['Normal'].paragraph_format.space_after=Pt(7)
d.styles['Title'].font.size=Pt(30); d.styles['Heading 1'].font.size=Pt(20); d.styles['Heading 2'].font.size=Pt(13)
def p(t): d.add_paragraph(t)
def h(t): d.add_heading(t,2)
def page(t): d.add_page_break(); d.add_heading(t,1)
def table(headers,rows,widths):
 t=d.add_table(rows=1,cols=len(headers)); t.autofit=False
 for c,w in zip(t.columns,widths): c.width=Inches(w)
 for c,v in zip(t.rows[0].cells,headers): c.text=v
 for row in rows:
  for c,v in zip(t.add_row().cells,row): c.text=v
 for i,r in enumerate(t.rows):
  for c,w in zip(r.cells,widths):
   c.width=Inches(w); pr=c._tc.get_or_add_tcPr(); sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),'19364A' if i==0 else ('F1F4F6' if i%2 else 'FFFFFF')); pr.append(sh)
   borders=OxmlElement('w:tcBorders')
   for edge in ['top','left','bottom','right']:
    e=OxmlElement('w:'+edge); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:color'),'D9D9D9'); borders.append(e)
   pr.append(borders)
   mar=OxmlElement('w:tcMar')
   for edge in ['top','left','bottom','right']:
    e=OxmlElement('w:'+edge); e.set(qn('w:w'),'95'); e.set(qn('w:type'),'dxa'); mar.append(e)
   pr.append(mar)
   for para in c.paragraphs:
    para.paragraph_format.space_after=Pt(3); para.paragraph_format.space_before=Pt(3)
    for run in para.runs:
     run.font.size=Pt(10)
     if i==0: run.bold=True; run.font.color.rgb=RGBColor(255,255,255)
  trpr=r._tr.get_or_add_trPr(); trpr.append(OxmlElement('w:cantSplit'))
 t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
 d.add_paragraph().paragraph_format.space_after=Pt(0)
d.add_heading('Saga Miners Boss Design',0)
d.add_paragraph('The Foreman and squad encounter rules',style='Subtitle')
p('Design proposal • Version 0.1 • October 2 2026')
p('Build the first boss around a corrupted mining exosuit called The Foreman. The encounter should turn guarding, dodging and elemental loadout choices into a coordinated squad challenge. This document defines the shared boss rules and a production specification for the first encounter.')
h('Required squad access')
p('Bosses require a queued squad of at least 3 players. Use 3 to 4 players for the first release, matching the existing party limit. There is no solo or duo entry. A public player cannot walk into an active boss arena or damage its boss.')
h('First boss identity')
table(['Attribute','Proposed direction'],[
('Name','The Foreman'),('Location','An abandoned extraction chamber reached from the Crystal Wilds'),('Fantasy','A mining supervisor trapped inside a crystal-corrupted industrial exosuit'),('Combat role','Slow, readable bruiser with squad positioning mechanics'),('Element profile','Thermal strength, Cryo weakness, neutral Ion and Corrosive'),('Target experience','A 3 to 5 minute first clear with an accessible, repeatable retry')],[1.35,5.6])
h('What the first encounter must teach')
p('Read the weapon before reacting. Guard the sweep, move out of the crusher, then spend resources while the machine cools. Coordinate during marked attacks instead of stacking on one teammate. Elemental weakness accelerates a clear but never becomes a required gem check.')
h('Existing game baseline')
p('The current enemy catalog contains Training Dummies, Crystal Crawlers and Shard Brutes. Brutus remains a neutral arena instructor. The Foreman, boss queue, encounter phases and resistance profiles below are proposed additions; this document does not change gameplay code.')
page('Squad queue and elemental rules')
h('Queue and encounter lifecycle')
p('The leader selects a boss and requests a ready check. Every queued member must be online, alive, outside another encounter and ready. Show the boss element profile before confirmation. Any roster or loadout change clears readiness. Revalidate the full roster on the server before entry and again before the encounter starts.')
p('Launch only when 3 or 4 valid members arrive. Lock the encounter roster and initial health scaling at the pull. Spectators and late joiners cannot contribute damage or receive encounter rewards. Before the pull, a drop below 3 cancels the start and returns the squad to staging.')
p('During combat, pause and protect all participants if connected encounter membership drops below 3. Allow a 60 second reconnect window for the same roster; resume only with at least 3 connected members. Otherwise reset and return the squad to staging. Defeated but connected members still count toward squad membership. A full wipe resets the boss and all hazards, with no entry fee for this first encounter.')
h('Every monster has its own profile')
p('Store strengths and weaknesses on each monster archetype, separately from the element of its attacks. Use the same readable profile in the bestiary, queue card and target inspection. Different species must differ in at least one affinity; avoid a single resistance shared by every enemy.')
table(['Monster','Strength','Weakness','Other elements'],[
('Crystal Crawler','Cryo 0.75x','Thermal 1.25x','Ion and Corrosive 1.00x'),('Shard Brute','Ion 0.75x','Corrosive 1.25x','Thermal and Cryo 1.00x'),('The Foreman','Thermal 0.75x','Cryo 1.25x','Ion and Corrosive 1.00x')],[1.55,1.4,1.45,2.55])
p('These are proposed starting values, not current combat values. Multipliers affect only the matching elemental damage component, including that element’s damage over time. Physical damage stays at 1.00x. Apply affinity once, after existing damage modifiers and before final rounding. Do not multiply healing, shield strength, resource drain or status duration through this table.')
h('Readable and fair affinities')
p('Use an element icon, name and explicit WEAK or RESISTS text. Keep the existing Thermal red, Cryo blue, Ion yellow and Corrosive green palette. Floating combat text can label the elemental result without recoloring physical damage. No elemental immunity or random resistance changes in the first boss. All armor families and gem combinations must remain viable.')
page('The Foreman appearance and arena')
h('Silhouette and materials')
p('Use a broad industrial torso, one oversized hydraulic crusher arm and a smaller gripping arm. A sunken helmet with a horizontal amber visor keeps the face readable. Broken shoulder plating exposes a pulsing reactor, while crystal growth runs through the back and one leg. Keep the silhouette mechanical and grounded in the existing mining equipment.')
p('Build the material hierarchy from dark steel structure, worn ochre safety panels, pale metal edges and localized Thermal glow. Reserve the brightest emissive surfaces for active vents and attack tells. Avoid a full glowing body that hides weapon motion. The boss should be about 1.7 times player height, with its upper body visible in first person at normal melee range.')
h('Animation direction')
p('Idle: a heavy weight shift, uneven piston motion and brief vent breaths. Sweep: the crusher draws clearly across the opposite shoulder. Crusher: the arm lifts overhead and visibly locks before dropping. Recovery: the weapon lodges in the floor and the reactor opens. Phase transition: a shoulder cover ejects and vents flare outward. Defeat: pressure releases, the knees buckle and the visor fades.')
h('Arena composition')
p('Prototype a circular chamber roughly 64 studs across, with a clear central fighting area at least 40 studs wide. Put extraction machinery and spectators outside the playable rim. Three nonblocking floor landmarks help teammates call positions. Avoid central pillars that interrupt first-person target lock or hide ground warnings.')
p('Use a dark matte floor with restrained grid lines. Telegraphs sit above the grid and use both animated boundaries and simple symbols. Keep the space around the player’s feet visible; ground danger must also have a directional HUD cue when its source is off-screen.')
h('Boss and squad interface')
p('Use one compact boss health bar at top center, with the boss name, phase label and affinity icons. Show a cast label only while an attack is being signaled. A narrow squad panel lists health, defeated state and reconnect status. Keep Health, Stamina and Charge unobstructed at bottom center. Never force the camera to face the boss or disable the player’s turn-away unlock.')
h('Required first pass assets')
table(['Discipline','Deliverables'],[
('Character art','Exosuit mesh or modular model, crusher, reactor, visor and damaged shoulder state'),('Animation','Idle, locomotion, sweep, crusher, vent, mark cast, stagger, transition and defeat'),('Effects and audio','Sweep trail, ground outline, impact debris, vent heat, target marker and distinct windup sounds'),('Interface','Boss bar, affinity icons, squad ready check, reconnect state and personal reward panel')],[1.3,5.65])
page('Attacks and encounter phases')
p('All timings and damage below are prototype tuning. Damage is expressed as a percentage of a reference starter character’s maximum health before mitigation; convert it to fixed encounter damage during implementation. Do not scale each hit to the victim’s health, which would erase the benefit of Health gear.')
table(['Attack','Signal and timing','Response and payoff'],[
('Piston Sweep\n12% damage','0.9 s shoulder windup; frontal arc; 1.2 s recovery.','Guard or step outside the arc. A well-timed guard uses the existing guard rules.'),
('Hydraulic Crusher\n22% damage','1.4 s overhead windup; marked 8 stud radius; 2.0 s recovery.','Dodge or walk out. Cannot be guarded. Aim stops tracking for the final 0.4 s.'),
('Pressure Vent\n8% per pulse','1.5 s warning; three clearly marked radial lanes; two pulses 1 s apart.','Move into the safe gaps. No damage between lanes; never cover the entire arena.'),
('Marked Extraction\n16% damage','Mark two different living players for 2.0 s; each leaves one 7 stud impact circle.','Spread the markers away from teammates, then leave the circles. No shared damage requirement.')],[1.35,2.75,2.85])
h('Phase one from full health to 60 percent')
p('Open with a Sweep after a 2 second grace period. Alternate Sweep and Crusher with a minimum 2 second movement interval after recovery. Introduce Pressure Vent only after the squad has seen both basic attacks. Use a fixed opening pattern so the first attempt teaches rather than surprises.')
h('Phase two below 60 percent')
p('Finish the current attack before changing phase. Use a 2 second non-damaging pressure release, then add Marked Extraction. Keep existing windup durations. Increase variety through the sequence, not faster unreadable animations. Never overlap Marked Extraction impacts with Pressure Vent. Maintain at least one safe route for every living player.')
h('Recovery and control')
p('After every Crusher, expose the reactor for 2 seconds as a clear attack opportunity; hits still use the same affinity table. For this first prototype, exposure has no extra damage multiplier. Boss stun effects contribute to a stagger meter instead of repeatedly freezing its AI. Start with one 2 second stagger per 20 seconds maximum; other crowd-control rules require explicit boss-specific tuning.')
h('Target selection')
p('Follow the existing threat and taunt integration. Telegraphs lock to a target at cast start and stop tracking at the stated cutoff. If that target dies, cancel the cast rather than snap its impact onto another teammate. Do not begin another attack until the current recovery has ended.')
page('Rewards implementation and review')
h('Squad balance and rewards')
p('Start with one 3-player health budget that produces a 3 to 5 minute clear using ordinary starter equipment. For 4 players, prototype 1.30 times that health with unchanged damage and timings. Fix the scaling when the fight begins; do not reduce boss health when a player disconnects.')
p('On a clear, give each eligible roster member a personal reward: one armor item, one gem and coins. Roll armor family uniformly, then slot uniformly; use a separate first-boss rarity table. Prototype 70% Uncommon and 30% Rare armor, with no Ultra Rare or Legendary drops. Gem rarity follows its own supported element rules. Quantity, coin value and repeat-clear economy remain tuning decisions.')
p('Give a guaranteed cosmetic Foreman insignia on the first clear, with no combat bonus. Never require top damage for credit. Eligibility requires joining the pull and participating through damage, guarding, support or encounter actions. Defeated contributors remain eligible. Record each reward by encounter ID and player ID so retries and reconnects cannot duplicate it; reserve delivery for a full inventory.')
h('Engineering sequence')
p('1. Add shared BossDefinitions and monster affinity profiles. Reuse internal Fire, Frost, Shock and Poison keys while displaying Thermal, Cryo, Ion and Corrosive.')
p('2. Add a server-owned boss queue and encounter roster around the existing Party system. Implement ready checks, entry validation, reconnect pause, wipe reset and cleanup before rewards.')
p('3. Extend enemy behavior with Telegraph, Execute, Recover, Transition and Defeated states. Keep damage, target eligibility, hit volumes, phase changes and reward grants authoritative on the server.')
p('4. Deliver the graybox Foreman and all four attacks, then connect first-person cues and boss UI. Add final animation, materials, sound and effects after the fight is readable.')
h('Acceptance tests')
p('Reject solo and duo entry, forged queue requests and a roster falling below 3 before pull. Verify 3- and 4-player ready checks, reconnect pause and timeout reset. Outsiders cannot damage the boss or claim rewards.')
p('Check all four elements against every profile, mixed physical and elemental hits, and damage over time. Verify affinities apply once, guardable versus unguardable attacks behave correctly, and no attack can hit twice from duplicate effect messages.')
p('Complete the encounter with each armor family and without Cryo gems. Test low-health phase crossing, target death during windup, a squad wipe, full reward inventory and duplicate claim attempts. Review telegraph legibility on touch screens and with reduced effects enabled.')
h('Next design decisions')
p('Confirm The Foreman’s identity and industrial silhouette, then choose whether the second boss should contrast it with mobility or ranged pressure. Playtest the graybox before locking health, coin rewards, stagger thresholds or final art production.')
p('Implementation references: src/server/Systems/Enemies.luau; src/server/Systems/Party.luau; src/shared/GemConfig.luau; docs/target-lock.md. Boss rules and encounter specifications in this document are proposals except the required minimum squad of 3.')
d.save(out/'Saga-Miners-Boss-Design.docx')
print(out/'Saga-Miners-Boss-Design.docx')
